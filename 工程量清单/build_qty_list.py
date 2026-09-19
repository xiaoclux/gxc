# -*- coding: utf-8 -*-
"""从广联达导出的构件计算书 / 钢筋明细表 / 清单模板生成《分部分项工程量清单计算书》与《分部分项工程量清单》Excel。"""
import re, collections, sys
import xlrd, openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

UP = sys.argv[2] if len(sys.argv) > 2 else "./源数据/"
F_CALC = UP + "3c2df904-________.xls"
F_REBAR = UP + "93d01613-_____.xls"
F_LIST = UP + "4c1dcb22-__.xlsx"
OUT = sys.argv[1] if len(sys.argv) > 1 else "工程量清单及计算书.xlsx"

FLOOR_ORDER = ["基础层", "首层", "第2层", "第3层", "第4层", "第5层"]

# ---------------------------------------------------------------- 1. 解析构件计算书
# elems: list of dict(floor, typ, name, no, pos, items=[(label, expr_clean, value, unit)])
elems = []
wb = xlrd.open_workbook(F_CALC)
for sh in wb.sheets():
    floor = sh.name.split("-")[-1]
    cur = None
    for r in range(sh.nrows):
        a = str(sh.cell_value(r, 0)).strip()
        b = str(sh.cell_value(r, 1)).strip()
        c = str(sh.cell_value(r, 2)).strip()
        if a and a != "序号" and not re.match(r"^[\d.]+$", a):
            _, typ, name = [x.strip() for x in a.split(" - ", 2)]
            cur = (typ, name)
            continue
        if cur is None or "=" not in c:
            continue
        items = []
        for line in c.split("\n"):
            m = re.match(r"^(.+?)\s*=\s*(.*)=\s*([\d.]+)\s*(m3|m2|m)?\s*$", line.strip())
            if not m:
                continue
            label, expr, val, unit = m.group(1).strip(), m.group(2).strip(), float(m.group(3)), m.group(4) or ""
            expr = re.sub(r"<[^>]*>", "", expr).replace("*", "×").strip()
            items.append((label, expr, val, unit))
        elems.append(dict(floor=floor, typ=cur[0], name=cur[1], no=a, pos=b, items=items))

def sel(typ=None, name_re=None, floors=None, label=None, name_fn=None):
    """选出满足条件的 (floor, name, expr, value, unit)。"""
    out = []
    for e in elems:
        if typ and e["typ"] not in (typ if isinstance(typ, (list, tuple)) else [typ]):
            continue
        if floors and e["floor"] not in floors:
            continue
        if name_re and not re.search(name_re, e["name"]):
            continue
        if name_fn and not name_fn(e["name"]):
            continue
        for (lab, expr, val, unit) in e["items"]:
            if label and lab != label:
                continue
            out.append((e["floor"], e["name"], expr, val, unit))
    return out

def sel_room(sub_label, room_re=None, floors=None):
    """房间类构件：按子项名称（地面积/踢脚抹灰长度/墙面抹灰面积/天棚抹灰面积/墙裙块料面积）取。"""
    out = []
    for e in elems:
        if e["typ"] != "房间":
            continue
        if floors and e["floor"] not in floors:
            continue
        if room_re and not re.search(room_re, e["name"]):
            continue
        for (lab, expr, val, unit) in e["items"]:
            if lab == sub_label:
                out.append((e["floor"], e["name"], expr, val, unit))
    return out

# ---------------------------------------------------------------- 2. 解析钢筋明细表
# rebar[(floor, comp_base_name, diameter)] = [weights...]
rebar = collections.defaultdict(list)
wb2 = xlrd.open_workbook(F_REBAR)
s2 = wb2.sheet_by_index(0)
floor = comp = None
for r in range(s2.nrows):
    a = str(s2.cell_value(r, 0)).strip()
    if a.startswith("楼层名称"):
        floor = re.sub(r"楼层名称：|（.*", "", a)
        continue
    if a.startswith("构件名称"):
        comp = re.sub(r"\[.*", "", a.replace("构件名称：", ""))
        continue
    if a in ("筋号", "") or a.startswith("构件位置"):
        continue
    d = str(s2.cell_value(r, 2)).strip()
    d = str(int(float(d)))
    rebar[(floor, comp, d)].append(float(s2.cell_value(r, 9)))
REBAR_TOTAL = sum(sum(v) for v in rebar.values())

def rcat(n):
    if re.match(r"^\d+$", n): return "板"
    if n.startswith(("1KZ", "TZ")): return "柱"
    if n.startswith("GZ"): return "构造柱"
    if n.startswith(("DKL", "DL")): return "基础梁"
    if n.startswith("J-"): return "独基"
    if n.startswith("GL"): return "过梁"
    if n.startswith("QL"): return "圈梁"
    if n.startswith("LT"): return "楼梯"
    if n.startswith("YPL"): return "雨蓬梁"
    if n.startswith("YP"): return "雨蓬栏板"
    if re.match(r"^\d?(W?KL|W?L)", n): return "梁"
    raise ValueError(n)

def sel_rebar(cats, diam):
    out = []
    for (fl, comp, d), ws in rebar.items():
        if d == str(diam) and rcat(comp) in cats:
            out.append((fl, comp, ws))
    return out

# ---------------------------------------------------------------- 3. 读清单模板
wb3 = openpyxl.load_workbook(F_LIST)
ws3 = wb3["Sheet1"]
qlist = []  # dict(code,name,feat,unit)
for row in ws3.iter_rows(min_row=2, values_only=True):
    if row[0]:
        qlist.append(dict(code=str(row[0]).strip(), name=str(row[1]).strip(), feat=str(row[2]).strip(), unit=str(row[3]).strip()))

# ---------------------------------------------------------------- 4. 清单项 → 数据源 映射
# 每项返回 rows: list of (计算部位, 计算式, 结果, 单位) ；以及 备注
def fmt(v, nd=4):
    s = f"{v:.{nd}f}".rstrip("0").rstrip(".")
    return s if s else "0"

def group_rows(sel_rows, unit_out=None, part_fn=None):
    """把 (floor,name,expr,val,unit) 按 (floor,name) 归并；相同表达式合并为 (expr)×n。"""
    byk = collections.OrderedDict()
    for fl, nm, expr, val, unit in sel_rows:
        k = (fl, nm)
        byk.setdefault(k, []).append((expr, val, unit))
    rows = []
    for (fl, nm), lst in sorted(byk.items(), key=lambda kv: (FLOOR_ORDER.index(kv[0][0]), kv[0][1])):
        cnt = collections.OrderedDict()
        for expr, val, unit in lst:
            cnt.setdefault(expr, [0, 0.0, unit])
            cnt[expr][0] += 1
            cnt[expr][1] += val
        parts = []
        for expr, (n, v, unit) in cnt.items():
            e = expr if n == 1 else f"[{expr}]×{n}"
            parts.append(e)
        total = sum(v for (_, v, _) in lst)
        unit = unit_out or lst[0][2]
        part = part_fn(fl, nm) if part_fn else f"{fl} {nm}"
        rows.append((part, " + ".join(parts) + f" = {fmt(total)}", total, unit))
    return rows

def rebar_rows(cats, diam):
    byk = collections.OrderedDict()
    for fl, comp, ws in sel_rebar(cats, diam):
        byk.setdefault((fl, comp), []).extend(ws)
    rows = []
    for (fl, comp), ws in sorted(byk.items(), key=lambda kv: (FLOOR_ORDER.index(kv[0][0]), kv[0][1])):
        tot = sum(ws)
        rows.append((f"{fl} {comp} Φ{diam}", " + ".join(fmt(w, 3) for w in ws) + f" = {fmt(tot, 3)}", tot, "kg"))
    return rows

def const_row(part, expr, val, unit):
    return [(part, expr, val, unit)]

# 分类选择器
COL = sel(typ="柱", label="体积")
GZ = sel(typ="构造柱", label="体积")
BEAM = sel(typ="梁", label="体积", name_fn=lambda n: not n.startswith("YPL"))
YPL = sel(typ="梁", label="体积", name_fn=lambda n: n.startswith("YPL"))
SLAB = sel(typ="现浇板", label="体积")
LT = sel(typ="楼梯", label="体积")
QL = sel(typ="圈梁", label="体积")
GL = sel(typ="过梁", label="体积")
JC = sel(typ="独立基础", label="体积")
DC = sel(typ="垫层", label="体积")
WA = sel(typ="砌体墙", label="体积", name_re=r"WQ")
WB_ = sel(typ="砌体墙", label="体积", name_re=r"NQ")
WN = sel(typ="砌体墙", label="体积", name_re=r"女儿墙")
TF = sel(typ="大开挖土方", label="土方体积")
HT = sel(typ="大开挖土方", label="素土回填体积")
JZ = sel(typ="建筑面积", label="面积", floors=["首层"])
SS = sel(typ="散水", label="面积")
TJ = sel(typ="台阶", label="台阶整体水平投影面积")
YD = sel(typ="压顶", label="体积")
YP = sel(typ="栏板", label="体积")
DOOR = sel(typ="门", label="洞口面积")
WIN = sel(typ="窗", label="洞口面积")
WM_FS = sel(typ="屋面", label="防水面积")
WM_JB = sel(typ="屋面", label="卷边面积")
EXT = sel(typ="墙面", label="墙面抹灰面积", name_re=r"外墙面")
INT5 = sel(typ="墙面", label="墙面抹灰面积", name_re=r"内墙面")
R_FLOOR_OTHER = sel_room("地面积", room_re=r"其他")
R_FLOOR_WC = sel_room("地面积", room_re=r"卫生间")
R_FLOOR_1F = sel_room("地面积", floors=["首层"])
R_TJ = sel_room("踢脚抹灰长度", room_re=r"其他")
R_WALL = sel_room("墙面抹灰面积")
R_QQ = sel_room("墙裙块料面积")
R_TP = sel_room("天棚抹灰面积")

sum_v = lambda rows: sum(x[3] for x in rows)
TF_T, HT_T = sum_v(TF), sum_v(HT)
A_1F = sum_v(R_FLOOR_1F)
JZ_A = sum_v(JZ)
FS_A, JB_A = sum_v(WM_FS), sum_v(WM_JB)
FLAT_A = FS_A - JB_A

NOTES = {}
def note(code, txt):
    NOTES.setdefault(code, []).append(txt)

def rows_for(item):
    code, name, feat = item["code"], item["name"], item["feat"]
    # ---------- 土石方
    if code == "010102001001":
        return group_rows(TF, part_fn=lambda f, n: f"{f} {n}(挖方)")
    if code == "010102007001":
        note(code, "按广联达大开挖土方的“素土回填体积”（基坑回填）汇总；房心回填未含。")
        return group_rows(HT, part_fn=lambda f, n: f"{f} {n}(回填)")
    if code == "010103001001":
        note(code, "计算书中无平整场地构件，按首层建筑面积计取；清单模板单位为 m³，规范应为 m²，请核对。")
        return const_row("首层 建筑面积 JZMJ-1", f"{fmt(JZ_A)} = {fmt(JZ_A)}", JZ_A, "m2")
    if code == "010103002001":
        note(code, "余方弃置 = 挖方总量 − 基坑回填总量（推导值）。")
        return const_row("基础层 DKW-1", f"{fmt(TF_T)} − {fmt(HT_T)} = {fmt(TF_T-HT_T)}", TF_T - HT_T, "m3")
    # ---------- 砌体
    if code == "010401003001": return group_rows(WA)
    if code == "010401003002": return group_rows(WB_)
    if code == "010401003003": return group_rows(WN)
    # ---------- 垫层
    if code == "010501001001": return group_rows(DC)
    if code == "010501002001":
        note(code, "计算书中无楼地面垫层构件，按首层各房间地面积×0.10m 推导。")
        return const_row("首层 地面(其他+卫生间)", f"{fmt(A_1F)}×0.1 = {fmt(A_1F*0.1)}", A_1F * 0.1, "m3")
    if code == "010501002002":
        note(code, "计算书中无楼地面垫层构件，按首层各房间地面积×0.15m 推导。")
        return const_row("首层 地面(其他+卫生间)", f"{fmt(A_1F)}×0.15 = {fmt(A_1F*0.15)}", A_1F * 0.15, "m3")
    # ---------- 混凝土
    if code == "010502001001": return group_rows(JC)
    if code == "010502006001":
        note(code, "含各层框架柱 1KZ1~1KZ4 及梯柱 TZ；基础层柱为基础顶至±0.000 段。")
        return group_rows(COL)
    if code == "010502011001":
        note(code, "含基础梁 DKL/DL、各层框架梁 KL/WKL、次梁 L/WL；雨棚梁 YPL 计入“其他板(雨棚板)”。")
        return group_rows(BEAM)
    if code == "010502013001": return group_rows(SLAB)
    if code == "010502019001":
        note(code, "广联达模型中未绘制雨棚板本体，仅有雨棚梁 YPL 与雨棚栏板 YP，此处为两者体积之和；雨棚板本体需补充。")
        return group_rows(YPL + YP)
    if code == "010502020001":
        note(code, "楼梯按广联达“梯段+梯梁+平台板”体积汇总（清单模板单位为 m³）。")
        return group_rows(LT)
    if code == "010502021001": return group_rows(GZ)
    if code == "010502022001": return group_rows(QL)
    if code == "010502023001": return group_rows(GL)
    if code == "010502029001": return group_rows(SS)
    if code == "010502032001": return group_rows(TJ)
    if code == "010503016001": return group_rows(YD)
    # ---------- 钢筋
    if code == "010506001001":
        note(code, "基础钢筋按独立基础 J-1/J-2/J-3 汇总（均为Φ12）；基础梁 DKL/DL 钢筋计入“现浇混凝土梁钢筋”对应直径。")
        return rebar_rows(["独基"], 12)
    if code.startswith("010506002"):
        d = int(re.search(r"直径(\d+)", feat).group(1))
        rws = rebar_rows(["柱"], d)
        if not rws:
            note(code, f"钢筋明细表中柱构件无Φ{d}钢筋，工程量为 0。")
        return rws
    if code.startswith("010506005"):
        d = int(re.search(r"直径(\d+)", feat).group(1))
        if d == 6 or d == 14:
            note(code, "含雨棚梁 YPL 钢筋（清单模板“雨蓬”项仅列Φ8）。")
        return rebar_rows(["梁", "基础梁", "雨蓬梁"], d)
    if code.startswith("010506006"):
        d = int(re.search(r"直径(\d+)", feat).group(1))
        return rebar_rows(["板"], d)
    if code.startswith("010506009"):
        d = int(re.search(r"直径(\d+)", feat).group(1))
        return rebar_rows(["楼梯"], d)
    if code.startswith("010506010"):
        d = int(re.search(r"直径(\d+)", feat).group(1))
        part = feat.split("\n")[0].replace("1.", "")
        cats = {"构造柱": ["构造柱"], "过梁": ["过梁"], "圈梁": ["圈梁"], "雨蓬": ["雨蓬栏板"]}[part]
        return rebar_rows(cats, d)
    if code == "010506020001":
        note(code, "钢筋明细表中无屋面网片；按刚性层面积×Φ4@150双向估算：Φ4 单重 0.099kg/m，每 m² 用量 2/0.15×0.099≈1.316kg/m²。")
        w = FLAT_A * 1.316
        return const_row("第5层 屋面 WM-1", f"{fmt(FLAT_A)}×1.316 = {fmt(w,3)}", w, "kg")
    # ---------- 门窗
    if code == "010801001001": return group_rows(DOOR)
    if code == "010807001001": return group_rows(WIN)
    # ---------- 屋面
    if code == "010902001001":
        note(code, "按广联达屋面“防水面积”（含 250 高卷边）汇总。")
        return group_rows(WM_FS)
    if code == "010902004001":
        note(code, "刚性层 = 屋面防水面积 − 卷边面积（水平投影）。")
        return const_row("第5层 屋面 WM-1", f"{fmt(FS_A)} − {fmt(JB_A)} = {fmt(FLAT_A)}", FLAT_A, "m2")
    if code == "010902005001":
        note(code, "计算书与钢筋表中均无雨水管信息，无法计取，需按图纸补充（根数×长度）。")
        return None
    if code == "010904002001":
        note(code, "按首层卫生间地面积计取，未含上翻高度部分。")
        return group_rows(R_FLOOR_WC)
    if code == "011001001001":
        note(code, "保温层 = 屋面防水面积 − 卷边面积（水平投影）。")
        return const_row("第5层 屋面 WM-1", f"{fmt(FS_A)} − {fmt(JB_A)} = {fmt(FLAT_A)}", FLAT_A, "m2")
    # ---------- 楼地面
    if code == "011101001001":
        note(code, "按各层“其他”房间地面积汇总（首层地面一并计入）。")
        return group_rows(R_FLOOR_OTHER)
    if code == "011102003001": return group_rows(R_FLOOR_WC)
    if code == "011105001001":
        note(code, "仅计“其他”房间踢脚；卫生间 2.1m 以下为瓷砖墙裙，不计踢脚。")
        return group_rows(R_TJ)
    # ---------- 墙柱面
    if code == "011201001001": return group_rows(EXT)
    if code == "011201001002":
        note(code, "含各层房间内墙面（卫生间为墙裙以上部分）及第5层女儿墙内侧/楼梯间墙面。")
        return group_rows(R_WALL + INT5)
    if code == "011203003001": return group_rows(R_QQ)
    if code == "011301001001": return group_rows(R_TP)
    if code == "011404001001":
        note(code, "外墙涂料面积 = 外墙一般抹灰面积。")
        return group_rows(EXT)
    if code == "011404001002":
        note(code, "内墙涂料面积 = 内墙一般抹灰面积。")
        return group_rows(R_WALL + INT5)
    if code == "011404002001":
        note(code, "天棚涂料面积 = 天棚抹灰面积。")
        return group_rows(R_TP)
    raise ValueError("未映射的清单项: " + code)

# ---------------------------------------------------------------- 5. 计算
results = []  # per item: dict(item, rows, total, qty, unit_out)
used_rebar = 0.0
for it in qlist:
    rws = rows_for(it)
    if rws is None:
        results.append(dict(item=it, rows=[], total=None, qty=None))
        continue
    total = sum(r[2] for r in rws)
    if it["unit"] == "t":
        qty = round(total / 1000.0, 3)
        if not it["code"].startswith("010506020"):
            used_rebar += total
    else:
        qty = round(total, 2)
    results.append(dict(item=it, rows=rws, total=total, qty=qty))

# 钢筋覆盖性核对
assert abs(used_rebar - REBAR_TOTAL) < 0.01, (used_rebar, REBAR_TOTAL)

# ---------------------------------------------------------------- 6. 写 Excel
thin = Side(style="thin", color="000000")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
F = Font(name="宋体", size=10)
FB = Font(name="宋体", size=10, bold=True)
FT = Font(name="宋体", size=14, bold=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=True)
FILL = PatternFill("solid", fgColor="F2F2F2")

out = openpyxl.Workbook()

# ---- Sheet 1 计算书
ws = out.active
ws.title = "分部分项工程量清单计算书"
heads = ["序号", "项目名称", "计算部位", "单位", "计算式", "结果"]
widths = [6, 16, 22, 6, 90, 12]
ws.merge_cells("A1:F1")
ws["A1"] = "分部分项工程量清单计算书"
ws["A1"].font = FT; ws["A1"].alignment = CENTER
ws.row_dimensions[1].height = 28
for i, (h, w) in enumerate(zip(heads, widths), 1):
    c = ws.cell(row=2, column=i, value=h); c.font = FB; c.alignment = CENTER; c.border = BORDER
    ws.column_dimensions[get_column_letter(i)].width = w
r = 3
for idx, res in enumerate(results, 1):
    it = res["item"]; rws = res["rows"]
    start = r
    unit_disp = it["unit"]
    if not rws:
        ws.cell(row=r, column=1, value=idx)
        ws.cell(row=r, column=2, value=it["name"])
        ws.cell(row=r, column=3, value="—")
        ws.cell(row=r, column=4, value=unit_disp)
        ws.cell(row=r, column=5, value="计算书中无依据，需补充")
        ws.cell(row=r, column=6, value="需补充")
        r += 1
    else:
        for (part, expr, val, u) in rws:
            ws.cell(row=r, column=3, value=part)
            ws.cell(row=r, column=4, value={"m3": "m³", "m2": "㎡", "m": "m", "kg": "kg"}.get(u, u))
            ws.cell(row=r, column=5, value=expr)
            r += 1
        # 汇总行：不写“小计”字样，其余内容保留
        ws.cell(row=r, column=4, value=unit_disp)
        if it["unit"] == "t":
            ws.cell(row=r, column=5, value=f"({' + '.join(fmt(x[2],3) for x in rws)}) ÷ 1000 = {res['qty']:.3f}")
        elif len(rws) == 1:
            ws.cell(row=r, column=5, value=f"{fmt(rws[0][2])} ≈ {res['qty']:.2f}")
        else:
            ws.cell(row=r, column=5, value=f"{' + '.join(fmt(x[2]) for x in rws)} = {fmt(res['total'])} ≈ {res['qty']:.2f}")
        ws.cell(row=r, column=6, value=res["qty"])
        for col in range(1, 7):
            ws.cell(row=r, column=col).fill = FILL
            ws.cell(row=r, column=col).font = FB
        sub = r          # 小计行（结果单元格单独保留，不并入上方合并区）
        r += 1
        ws.cell(row=start, column=1, value=idx)
        ws.cell(row=start, column=2, value=it["name"])
        ws.cell(row=start, column=6, value=res["qty"])
        ws.cell(row=start, column=6).font = FB
        if r - 1 > start:
            ws.merge_cells(start_row=start, start_column=1, end_row=r - 1, end_column=1)
            ws.merge_cells(start_row=start, start_column=2, end_row=r - 1, end_column=2)
        if sub - 1 > start:
            ws.merge_cells(start_row=start, start_column=6, end_row=sub - 1, end_column=6)
    for rr in range(start, r):
        txt = str(ws.cell(row=rr, column=5).value or "")
        lines = max(1, -(-len(txt.encode("gbk", "replace")) // 170))
        ws.row_dimensions[rr].height = 14 * lines + 4
        for col in range(1, 7):
            c = ws.cell(row=rr, column=col)
            c.border = BORDER
            if c.font != FB: c.font = F
            c.alignment = LEFT if col in (3, 5) else CENTER
            if col == 6:
                c.number_format = "0.000" if it["unit"] == "t" else "0.00"
                c.alignment = Alignment(horizontal="center", vertical="center")
ws.freeze_panes = "A3"
ws.print_title_rows = "1:2"
ws.page_setup.orientation = "landscape"; ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.oddFooter.center.text = "第 &P 页 共 &N 页"
ws.print_options.horizontalCentered = True

# ---- Sheet 2 清单
w2 = out.create_sheet("分部分项工程量清单")
heads2 = ["序号", "项目编码", "项目名称", "项目特征描述", "单位", "工程量"]
widths2 = [6, 15, 18, 60, 7, 12]
w2.merge_cells("A1:F1")
w2["A1"] = "分部分项工程量清单"; w2["A1"].font = FT; w2["A1"].alignment = CENTER
w2.row_dimensions[1].height = 28
for i, (h, w) in enumerate(zip(heads2, widths2), 1):
    c = w2.cell(row=2, column=i, value=h); c.font = FB; c.alignment = CENTER; c.border = BORDER
    w2.column_dimensions[get_column_letter(i)].width = w
r = 3
for idx, res in enumerate(results, 1):
    it = res["item"]
    vals = [idx, it["code"], it["name"], it["feat"], it["unit"], res["qty"] if res["qty"] is not None else "需补充"]
    for col, v in enumerate(vals, 1):
        c = w2.cell(row=r, column=col, value=v)
        c.font = F; c.border = BORDER
        c.alignment = LEFT if col == 4 else CENTER
        if col == 6 and isinstance(v, float):
            c.number_format = "0.000" if it["unit"] == "t" else "0.00"
    w2.row_dimensions[r].height = max(15, 14 * (it["feat"].count("\n") + 1) + 4)
    r += 1
w2.freeze_panes = "A3"
w2.print_title_rows = "1:2"
w2.page_setup.orientation = "portrait"; w2.page_setup.paperSize = w2.PAPERSIZE_A4
w2.page_setup.fitToWidth = 1; w2.page_setup.fitToHeight = 0; w2.sheet_properties.pageSetUpPr.fitToPage = True
w2.oddFooter.center.text = "第 &P 页 共 &N 页"
w2.print_options.horizontalCentered = True

# ---- Sheet 3 编制说明
w3 = out.create_sheet("编制说明")
w3.column_dimensions["A"].width = 14; w3.column_dimensions["B"].width = 22; w3.column_dimensions["C"].width = 100
w3["A1"] = "编制说明"; w3["A1"].font = FT
general = [
    "数据来源：①广联达《绘图输入构件工程量计算书》（基础层~第5层）；②广联达《钢筋明细表》；③清单模板（74 项）。",
    "计算书粒度：每层每个构件编号一行，同编号下计算式相同的图元合并为 [计算式]×n；每个清单项末行为小计。",
    "计算式已去掉广联达的尖括号注释，仅保留数值式；各图元结果取广联达原值（4 位小数），清单工程量 m³/㎡/m 取 2 位小数，钢筋 t 取 3 位小数。",
    "钢筋明细表导出的“级别”列为空，钢筋级别按清单模板“构件类型+直径→级别”对应关系归类；直径 12/14 过梁钢筋按模板标注 HPB300 保留。",
    f"钢筋校核：钢筋明细表总重 {REBAR_TOTAL:.3f}kg，已全部分配到各钢筋清单项（不含估算的屋面钢筋网片），分配合计 {used_rebar:.3f}kg。",
    "标注“需补充”的项目在计算书/钢筋表中无依据，需按图纸另行计算。",
]
rr = 2
for g in general:
    w3.cell(row=rr, column=1, value="总体"); w3.cell(row=rr, column=3, value=g)
    rr += 1
rr += 1
w3.cell(row=rr, column=1, value="项目编码").font = FB; w3.cell(row=rr, column=2, value="项目名称").font = FB; w3.cell(row=rr, column=3, value="说明").font = FB
rr += 1
name_of = {it["code"]: it["name"] for it in qlist}
for code, txts in NOTES.items():
    for t in txts:
        w3.cell(row=rr, column=1, value=code); w3.cell(row=rr, column=2, value=name_of[code]); w3.cell(row=rr, column=3, value=t)
        rr += 1
for row in w3.iter_rows(min_row=2, max_row=rr):
    for c in row:
        c.font = F if c.font != FB else FB; c.alignment = LEFT

out.save(OUT)

# ---- 控制台摘要
print(f"钢筋总重 {REBAR_TOTAL:.3f} kg, 已分配 {used_rebar:.3f} kg")
for idx, res in enumerate(results, 1):
    it = res["item"]
    print(f"{idx:2d} {it['code']} {it['name']:14s} {it['unit']:3s} {res['qty']!s:>10}  rows={len(res['rows'])}")
